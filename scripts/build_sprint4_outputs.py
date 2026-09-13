#!/usr/bin/env python3
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'outputs'/'sprint4_release'; OUT.mkdir(parents=True,exist_ok=True)
BLUE=RGBColor(25,65,105); MUTED=RGBColor(96,116,135); LIGHT=RGBColor(235,243,250)
chapters=[
('第二章｜行业发展历程、驱动因素与资本流向','行业加速，规模商业化仍未被证明'),
('第三章｜技术体系与技术成熟度','部分任务进入现场，长期可靠性仍待证明'),
('第四章｜产业链、供应链、制造体系与价值捕获','试制条件已形成，可重复交付仍是关键'),
('第五章｜市场规模、竞争格局与商业化空间','市场现实是离散交易，需求兑现比远期预测更重要'),
('第六章｜应用场景、客户任务与真实需求','优先看具体任务，而不是行业标签'),
('第七章｜商业化成熟度','交易、交付、运行和复制是不同证据阶段'),
('第八章｜代表企业与产品商业化比较','数量不可直接横排，证据性质必须拆开'),
('第九章｜资源能力与竞争优势','资源只有转化为客户结果才形成优势'),
('第十章｜客户、开发者与公开用户反馈','公开反馈仍稀缺，使用摩擦需要持续验证'),
('第十一章｜客户与供应商经济性','价格、收入和现金流不能直接等同于单位经济性'),
('第十二章｜风险与边界','风险必须有作用机制，证据缺口不自动等于风险'),
('第十三章｜投资结论','保留期权，但等待交易与复购证据'),
('第十四章｜AI辅助研究方法与可复用SOP','AI提高检索与审计效率，不替代判断')]

def add_text(slide,text,x,y,w,h,size=24,color=BLUE,bold=False):
 box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.text=text; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color; return box

def build_ppt():
 prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
 blank=prs.slide_layouts[6]
 def slide(title,subtitle=None):
  s=prs.slides.add_slide(blank); add_text(s,title,.65,.45,12,0.6,26,BLUE,True)
  if subtitle: add_text(s,subtitle,.68,1.1,12,.45,13,MUTED); return s
  return s
 s=prs.slides.add_slide(blank); add_text(s,'中国具身智能机器人产业与代表企业商业化研究',.8,2.2,11.8,1.1,30,BLUE,True); add_text(s,'管理层简报｜闫纾慈　独立研究｜2026年9月11日',.85,3.55,11,.4,15,MUTED)
 s=slide('本报告的核心判断','市场正在从技术展示走向任务验证与早期交易，但统一、可复制的规模商业化尚未被公开证据证明。')
 for i,(h,j) in enumerate([('现实市场','工业机器人已有可统计规模，人形机器人仍以离散交易和项目证据为主。'),('需求入口','制造、物流和平台采购出现真实事件，但交易性质并不相同。'),('竞争结构','人形机器人进入的是已有方案覆盖的预算，而非空白市场。'),('投资含义','重点跟踪交付、正常运行、复购和跨客户复制。')]):
  x=.8+(i%2)*6.1; y=2+(i//2)*1.8; sh=s.shapes.add_shape(1,Inches(x),Inches(y),Inches(5.4),Inches(1.25)); sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT; sh.line.color.rgb=RGBColor(160,190,210); add_text(s,h,x+.25,y+.18,4.9,.3,17,BLUE,True); add_text(s,j,x+.25,y+.55,4.9,.5,12,RGBColor(35,45,55))
 s=slide('研究结论如何逐步收敛','从定义、任务、竞争到交易证据，章节结论采用同一条证据纪律。')
 for i,t in enumerate(['定义口径','现实任务','现有方案','交易与运行','复制与投资']):
  x=.8+i*2.45; sh=s.shapes.add_shape(5,Inches(x),Inches(2.7),Inches(2.0),Inches(1.0)); sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT; sh.line.color.rgb=RGBColor(160,190,210); add_text(s,t,x+.18,3.02,1.65,.3,14,BLUE,True)
  if i<4: add_text(s,'→',x+2.05,3.0,.3,.3,18,MUTED,True)
 s=slide('章节地图：从行业现实走向投资判断')
 for i,(h,j) in enumerate(chapters):
  x=.75+(i%2)*6.1; y=1.55+(i//2)*.72; add_text(s,h,x,y,5.6,.25,11,BLUE,True); add_text(s,j,x+4.0,y,2.0,.25,9,MUTED)
 s=slide('市场现实与需求入口','统一市场规模尚未形成，最可观察的是具体交易和任务验证。')
 add_text(s,'工业机器人：全球安装量与运营存量可统计',.85,1.8,5.5,.5,18,BLUE,True); add_text(s,'人形机器人：公司收入、订单、交付和采购事件是离散锚点',.85,2.55,5.5,.5,18,BLUE,True); add_text(s,'远期预测：必须同时说明金额、年份、定义和范围，不能与现实交易相加。',.85,3.4,11,.55,16,MUTED)
 s=slide('代表企业比较：先区分证据性质')
 rows=[('宇树','产品销售/交付','规模信号更强，分部财务未知'),('智元','下线与交付','供给与交付证据，终端消化需继续验证'),('优必选','审计销量与人形收入','收入更可审计，集团与分部边界需区分'),('傅利叶','平台交付与康复渠道','平台交易真实，不等于生产替代')]
 table=s.shapes.add_table(len(rows)+1,3,Inches(.8),Inches(1.6),Inches(11.7),Inches(3.5)).table
 for c,t in enumerate(['企业','已观察证据','不能直接推出']): table.cell(0,c).text=t
 for r,row in enumerate(rows,1):
  for c,t in enumerate(row): table.cell(r,c).text=t
 for row in table.rows:
  for cell in row.cells:
   for p in cell.text_frame.paragraphs:
    p.font.name='Arial'; p.font.size=Pt(13); p.font.color.rgb=BLUE if row is table.rows[0] else RGBColor(35,45,55)
 s=slide('投资跟踪框架','把“潜力”转换成可观察的领先与滞后指标。')
 add_text(s,'领先指标',1.0,1.8,4,.4,20,BLUE,True); add_text(s,'付费采购、交付验收、正常运行、客户扩展',1.0,2.4,4.8,.8,18,RGBColor(35,45,55))
 add_text(s,'滞后指标',7.0,1.8,4,.4,20,BLUE,True); add_text(s,'复购、跨场地复制、分部收入、经营现金流改善',7.0,2.4,4.8,.8,18,RGBColor(35,45,55))
 s=slide('结论：保留期权，但等待可复制证据','当前最值得相信的是需求入口正在出现；最不应相信的是远期预测已经等于市场规模。')
 add_text(s,'下一步研究重点',.9,2,3.5,.4,20,BLUE,True); add_text(s,'交付是否进入正常运营\n运营是否降低人工介入\n客户是否持续复购\n供应商是否形成可持续现金流',.9,2.6,5,.9,18,RGBColor(35,45,55)); add_text(s,'本简报由正式研究报告转换而来，数字与来源应回溯至 Master HTML 的全局 Canonical Source Registry。',.9,5.9,11,.4,12,MUTED)
 path=OUT/'ReportFlow_Embodied_Intelligence_Executive_Brief.pptx'; prs.save(path); return path

def build_docx():
 d=Document(); sec=d.sections[0]; sec.top_margin=DInches(.7); sec.bottom_margin=DInches(.7); sec.left_margin=DInches(.85); sec.right_margin=DInches(.85)
 p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('中国具身智能机器人产业与代表企业商业化研究'); r.bold=True; r.font.size=DPt(22); r.font.name='Arial'
 p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('管理层简报｜闫纾慈　独立研究｜2026年9月11日'); r.font.size=DPt(11); r.font.color.rgb=__import__('docx').shared.RGBColor(96,116,135)
 d.add_heading('执行摘要',level=1); d.add_paragraph('市场正在从技术展示走向任务验证与早期交易，但统一、可复制的规模商业化尚未被公开证据证明。工业机器人已有可统计市场，人形机器人仍以离散交易、订单、交付和项目验证作为现实锚点。')
 d.add_heading('四项管理层结论',level=1)
 for t in ['现实市场：先区分可统计的工业机器人与离散的人形机器人交易。','需求入口：制造、物流和研发平台采购均可能是真实需求，但商业含义不同。','竞争结构：人形机器人面对的是工业机器人、机械臂、AMR、专机和人工等现有方案。','投资含义：跟踪交付、正常运行、复购与跨客户复制，不把远期预测当作已实现规模。']: d.add_paragraph(t,style='List Bullet')
 d.add_heading('代表企业证据对照',level=1); table=d.add_table(rows=1,cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,t in enumerate(['企业','已观察证据','边界']): table.rows[0].cells[i].text=t
 for row in [('宇树','产品销售与交付信号','分部财务与项目经济性未知'),('智元','下线与交付信号','终端消化与复购仍待验证'),('优必选','审计销量与人形收入','集团数据不能代理单台经济性'),('傅利叶','平台交付与康复渠道','平台交易不等于生产替代')]:
  cells=table.add_row().cells
  for i,t in enumerate(row): cells[i].text=t
 d.add_heading('跟踪指标',level=1); d.add_paragraph('领先指标：付费采购、交付验收、正常运行和客户扩展。\n滞后指标：复购、跨场地复制、分部收入与经营现金流改善。')
 d.add_heading('研究边界',level=1); d.add_paragraph('本简报是正式研究报告的输出转换，不新增研究结论。所有数字与引用应回溯至 Master HTML 的全局 Canonical Source Registry。')
 # Use a CJK-capable family for LibreOffice/Word rendering.
 for para in d.paragraphs:
  for run in para.runs: run.font.name='Hiragino Sans GB'
 for table in d.tables:
  for row in table.rows:
   for cell in row.cells:
    for para in cell.paragraphs:
     for run in para.runs: run.font.name='Hiragino Sans GB'
 path=OUT/'ReportFlow_Embodied_Intelligence_Executive_Brief.docx'; d.save(path); return path

if __name__=='__main__': print(build_ppt()); print(build_docx())
